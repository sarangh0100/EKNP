# src/extractors/parsers.py
import io
import fitz  # PyMuPDF
import docx  # python-docx
import pptx  # python-pptx
import openpyxl
from src.extractors.base import Extractor, ExtractionResult

class PdfExtractor(Extractor):
    """PDF content and metadata extraction via PyMuPDF (FR-E1)."""
    
    @property
    def supported_extensions(self):
        return ["pdf"]

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        res = ExtractionResult(extraction_method="PyMuPDF")
        try:
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            res.page_count = len(doc)
            res.native_metadata_raw = dict(doc.metadata) if doc.metadata else {}
            
            text_chunks = [page.get_text() for page in doc]
            res.text = "\n".join(text_chunks).strip()
            doc.close()
        except Exception as e:
            res.errors.append(f"PdfExtractor failed: {str(e)}")
        return res

class DocxExtractor(Extractor):
    """Word document content and core properties extractor (FR-E2)."""
    
    @property
    def supported_extensions(self):
        return ["docx"]

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        res = ExtractionResult(extraction_method="python-docx")
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            cp = doc.core_properties
            res.native_metadata_raw = {
                "title": cp.title or "",
                "author": cp.author or "",
                "last_modified_by": cp.last_modified_by or "",
                "created": str(cp.created) if cp.created else "",
                "modified": str(cp.modified) if cp.modified else "",
                "revision": cp.revision or 1,
            }
            res.text = "\n".join([p.text for p in doc.paragraphs if p.text]).strip()
        except Exception as e:
            res.errors.append(f"DocxExtractor failed: {str(e)}")
        return res

class PptxExtractor(Extractor):
    """PowerPoint presentation extractor (FR-E3)."""
    
    @property
    def supported_extensions(self):
        return ["pptx"]

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        res = ExtractionResult(extraction_method="python-pptx")
        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
            res.page_count = len(prs.slides)
            
            cp = prs.core_properties
            res.native_metadata_raw = {
                "title": cp.title or "",
                "author": cp.author or "",
                "last_modified_by": cp.last_modified_by or "",
                "created": str(cp.created) if cp.created else "",
                "modified": str(cp.modified) if cp.modified else "",
            }
            slide_texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        slide_texts.append(shape.text_frame.text)
            res.text = "\n".join(slide_texts).strip()
        except Exception as e:
            res.errors.append(f"PptxExtractor failed: {str(e)}")
        return res

class XlsxExtractor(Extractor):
    """Excel workbook properties and cell content extractor (FR-E4)."""
    
    @property
    def supported_extensions(self):
        return ["xlsx"]

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        res = ExtractionResult(extraction_method="openpyxl")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            res.page_count = len(wb.sheetnames)
            
            cp = wb.properties
            res.native_metadata_raw = {
                "title": cp.title or "",
                "creator": cp.creator or "",
                "last_modified_by": cp.lastModifiedBy or "",
                "created": str(cp.created) if cp.created else "",
                "modified": str(cp.modified) if cp.modified else "",
            }
            cell_data = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        cell_data.append(" | ".join(row_vals))
            res.text = "\n".join(cell_data).strip()
            wb.close()
        except Exception as e:
            res.errors.append(f"XlsxExtractor failed: {str(e)}")
        return res

class TxtExtractor(Extractor):
    """Plain text parser (FR-E5)."""
    
    @property
    def supported_extensions(self):
        return ["txt", "md"]

    def extract(self, file_bytes: bytes, filename: str) -> ExtractionResult:
        res = ExtractionResult(extraction_method="builtin_txt")
        try:
            res.text = file_bytes.decode("utf-8", errors="replace").strip()
            res.native_metadata_raw = {"encoding": "utf-8"}
            res.page_count = 1
        except Exception as e:
            res.errors.append(f"TxtExtractor failed: {str(e)}")
        return res