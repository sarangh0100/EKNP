# generate_dataset.py
import os
import io
import fitz        # PyMuPDF
import docx        # python-docx
import openpyxl
import pptx
from PIL import Image, ImageDraw, ImageFont

BASE_DIR = os.path.join("data", "silos")
SILOS = ["hr", "finance", "engineering", "legal"]

def create_dirs():
    for silo in SILOS:
        os.makedirs(os.path.join(BASE_DIR, silo), exist_ok=True)

def create_pdf(path, title, author, text):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), f"{title}\n\n{text}", fontsize=11)
    doc.set_metadata({"title": title, "author": author})
    doc.save(path)
    doc.close()

def create_docx(path, title, author, text):
    doc = docx.Document()
    cp = doc.core_properties
    cp.title = title
    cp.author = author
    doc.add_heading(title, level=1)
    doc.add_paragraph(text)
    doc.save(path)

def create_xlsx(path, title, author, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    wb.properties.title = title
    wb.properties.creator = author
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(path)

def create_pptx(path, title, author, text):
    prs = pptx.Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    prs.core_properties.title = title
    prs.core_properties.author = author
    slide.shapes.title.text = title
    slide.placeholders[1].text = text
    prs.save(path)

def create_image(path, text):
    img = Image.new("RGB", (600, 200), color=(255, 255, 255))
    d = ImageDraw.Draw(img)
    d.text((20, 80), text, fill=(0, 0, 0))
    img.save(path)

def create_txt(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)

def generate_all():
    create_dirs()
    print("Generating 30 files across silos...")

    # HR Silo (8 files)
    create_docx(os.path.join(BASE_DIR, "hr", "leave_policy_2024.docx"), "Annual Leave Policy 2024", "Jane Doe", "Full-time employees receive 25 paid vacation days annually.")
    create_pdf(os.path.join(BASE_DIR, "hr", "Remote_Work_Agreement_v2.pdf"), "Remote Work Guidelines", "DOE, JANE", "Eligibility criteria and technical requirements for remote employees.")
    create_pdf(os.path.join(BASE_DIR, "hr", "benefits_guide_final.pdf"), "Corporate Employee Benefits", "", "Comprehensive healthcare, retirement 401k match, and wellness stipends.")
    create_docx(os.path.join(BASE_DIR, "hr", "performance_review_template.docx"), "Quarterly Review Form", "HR Admin", "Section 1: Objective Key Results. Section 2: Leadership core competencies.")
    create_txt(os.path.join(BASE_DIR, "hr", "intern_onboarding_checklist.txt"), "1. Issue corporate laptop\n2. Configure email access\n3. Schedule mentor check-in.")
    create_image(os.path.join(BASE_DIR, "hr", "org_chart_q3.png"), "HR Department Organization Structure Q3 2024")
    create_docx(os.path.join(BASE_DIR, "hr", "emergency_contact_update.docx"), "Emergency Form", "smith, john", "Notify HR operations immediately upon changes to home address.")
    create_pdf(os.path.join(BASE_DIR, "hr", "scanned_memo.pdf"), "Notice to Employees", "Staff Office", "Office premises closed on national statutory holidays.")

    # Finance Silo (7 files)
    create_xlsx(os.path.join(BASE_DIR, "finance", "q3_budget_actuals.xlsx"), "Q3 Operating Budget", "Finance Dept", ["Cost Center", "Budget", "Actual"], [["Engineering", "500000", "482000"], ["HR", "120000", "118500"]])
    create_pdf(os.path.join(BASE_DIR, "finance", "travel_reimbursement_policy.pdf"), "Corporate Travel Policy", "CFO Office", "Receipts must be submitted within 30 days of trip conclusion.")
    create_xlsx(os.path.join(BASE_DIR, "finance", "vendor_invoices_oct2024.xlsx"), "Vendor Tracker", "Accounts Payable", ["Vendor", "Amount", "Status"], [["Cloud Services Inc", "4500", "Paid"], ["Office Supplies Co", "820", "Pending"]])
    create_docx(os.path.join(BASE_DIR, "finance", "audit_findings_2023.docx"), "Internal Audit Summary", "External Auditor", "All financial records adhere strictly to GAAP standards.")
    create_pptx(os.path.join(BASE_DIR, "finance", "fiscal_year_projections.pptx"), "FY25 Fiscal Forecast", "Treasury", "Anticipated 14% growth in ARR and gross margin expansion.")
    create_txt(os.path.join(BASE_DIR, "finance", "wire_transfer_limits.txt"), "Transactions exceeding $50,000 require dual signature executive authorization.")
    create_image(os.path.join(BASE_DIR, "finance", "receipt_sample.png"), "Expense Invoice Amount USD 142.50 Vendor Cafe")

    # Engineering Silo (8 files)
    create_docx(os.path.join(BASE_DIR, "engineering", "architecture_rfc_042.docx"), "RFC 042: Event Driven Architecture", "lead_architect@corp.local", "Proposing Kafka message brokers to decouple ingestion services.")
    create_pdf(os.path.join(BASE_DIR, "engineering", "api_guidelines_draft.pdf"), "REST API Standard v2", "API Working Group", "All public endpoints must accept and return application/json payloads.")
    create_txt(os.path.join(BASE_DIR, "engineering", "sprint_retro_notes.txt"), "What went well: CI build times dropped to 3 minutes. What to improve: Flaky integration tests.")
    create_pdf(os.path.join(BASE_DIR, "engineering", "disaster_recovery_plan.pdf"), "Disaster Recovery Runbook", "DevOps Team", "Target RPO is 15 minutes; target RTO is 1 hour across primary availability zones.")
    create_pptx(os.path.join(BASE_DIR, "engineering", "microservices_overview.pptx"), "Microservices Ecosystem", "Alice Wang", "High level map of customer facing APIs, authentication, and database tiers.")
    create_docx(os.path.join(BASE_DIR, "engineering", "db_migration_runbook.docx"), "PostgreSQL 16 Migration", "Wang, Alice", "Step-by-step procedures to upgrade database clusters and enable pgvector.")
    create_xlsx(os.path.join(BASE_DIR, "engineering", "latency_benchmarks_oct.xlsx"), "Latency Benchmarks", "SRE Team", ["Service", "p50_ms", "p99_ms"], [["Search API", "42", "110"], ["Ingest Worker", "180", "420"]])
    create_image(os.path.join(BASE_DIR, "engineering", "topology_diagram.png"), "Application Microservices Cluster Topology")

    # Legal Silo (7 files)
    create_docx(os.path.join(BASE_DIR, "legal", "nda_standard_template.docx"), "Mutual Non-Disclosure Agreement", "Legal Counsel", "Confidential information shall not be disclosed to any third parties without prior consent.")
    create_pdf(os.path.join(BASE_DIR, "legal", "data_privacy_addendum_gdpr.pdf"), "GDPR Data Processing Addendum", "Privacy Officer", "Customer data stored locally; no unauthorized cross-border transmissions.")
    create_xlsx(os.path.join(BASE_DIR, "legal", "patent_filing_status.xlsx"), "IP Docket Tracking", "Legal Dept", ["Application", "Jurisdiction", "Status"], [["Semantic Search Index", "USPTO", "Pending"], ["Vector Upsert Pipeline", "EPO", "Drafting"]])
    create_docx(os.path.join(BASE_DIR, "legal", "intellectual_property_policy.docx"), "Intellectual Property Policy", "General Counsel", "All inventions developed on corporate equipment remain company property.")
    create_txt(os.path.join(BASE_DIR, "legal", "regulatory_compliance_memo.txt"), "Annual reminder regarding insider trading blackouts and regulatory disclosure filing dates.")
    create_pdf(os.path.join(BASE_DIR, "legal", "litigation_hold_notice.pdf"), "Litigation Hold Instruction", "General Counsel", "Preserve all records and electronic files pertaining to Matter Ref 984.")
    # Intentionally corrupt file for FR-E8 verification
    with open(os.path.join(BASE_DIR, "legal", "corrupt_document.pdf"), "wb") as f:
        f.write(b"%PDF-1.4 CORRUPTED_BINARY_PAYLOAD_NOT_A_VALID_PDF")

    print("Successfully generated 30 test files with realistic, messy metadata!")

if __name__ == "__main__":
    generate_all()